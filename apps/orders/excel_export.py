"""
Lidlarni Excel (.xlsx) faylga eksport qilish.

Bu — loyihaning "ichki CRM emas" tamoyiliga mos, oddiy ko'chirib olish
vositasi: do'kon egasi hozirgi lidlar ro'yxatini (TWA'dagi filtrlar bilan
birga) bitta .xlsx fayl sifatida yuklab olib, o'zining Excel/CRM
jarayoniga qo'shishi mumkin.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("Lead raqami", "lead_number", 20),
    ("Sana", "created_at", 18),
    ("Ism-familiya", "customer_full_name", 25),
    ("Telefon", "customer_phone", 18),
    ("Qiziqqan kurs/mahsulot", "interested_course_name", 28),
    ("Holat", "lead_status_label", 14),
    ("Manba", "source_label", 16),
    ("Platforma", "platform_label", 14),
    ("Izoh", "notes", 40),
]


def _row_values(lead) -> dict:
    return {
        "lead_number": lead.lead_number,
        "created_at": lead.created_at.strftime("%Y-%m-%d %H:%M") if lead.created_at else "",
        "customer_full_name": lead.customer_full_name,
        "customer_phone": lead.customer_phone,
        "interested_course_name": lead.interested_course.name if lead.interested_course_id else "",
        "lead_status_label": lead.get_lead_status_display(),
        "source_label": lead.get_source_display(),
        "platform_label": lead.get_platform_display(),
        "notes": lead.notes or "",
    }


def build_leads_workbook(leads) -> BytesIO:
    """
    Berilgan Lead querysetidan (yoki ro'yxatidan) .xlsx workbook yaratadi.

    Returns:
        BytesIO — fayl kontenti (o'qish uchun tayyor, .seek(0) qilingan)
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Leadlar"

    header_font = Font(bold=True)
    for col_idx, (header, _key, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, lead in enumerate(leads, start=2):
        values = _row_values(lead)
        for col_idx, (_header, key, _width) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=values[key])
            # ⚠️ XAVFSIZLIK (formula injection): mijoz ismi/izohi tashqaridan keladi.
            # openpyxl "=" bilan boshlanadigan matnni avtomatik FORMULA deb saqlaydi
            # (masalan `=HYPERLINK("http://...","Bosing")`), ega faylni ochganda u ishga
            # tushadi. Hamma matnli katak majburan MATN turiga o'tkaziladi — Excel uni
            # hech qachon hisoblamaydi, faqat ko'rsatadi.
            if isinstance(values[key], str):
                cell.data_type = "s"

    ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
